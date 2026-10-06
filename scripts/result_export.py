from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence


TITLE_FILL = "1F4E78"
SECTION_FILL = "D9EAF7"
HEADER_FILL = "5B9BD5"
LIGHT_FILL = "F3F7FA"
ACCENT_FILL = "E2F0D9"
WARNING_FILL = "FFF2CC"
TEXT_COLOR = "1F2937"
WHITE = "FFFFFF"
FONT_NAME = "Microsoft YaHei"


def format_number(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if abs(value) < 5e-12:
            return "0"
        if abs(value) >= 1e5 or abs(value) < 1e-5:
            return f"{value:.6e}"
        return f"{value:.8f}".rstrip("0").rstrip(".")
    return str(value)


def _safe_name(text: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid else char for char in text)
    return cleaned.strip().strip(".") or "变焦计算结果"


def _make_output_dir(base_dir: Path, scheme: str) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidate = base_dir / f"{stamp}_{_safe_name(scheme)}"
    suffix = 1
    while candidate.exists():
        candidate = base_dir / f"{stamp}_{_safe_name(scheme)}_{suffix:02d}"
        suffix += 1
    candidate.mkdir(parents=True, exist_ok=False)
    return candidate


def _headers(rows: Sequence[Dict[str, Any]]) -> List[str]:
    result: List[str] = []
    for row in rows:
        for key in row:
            if key not in result:
                result.append(key)
    return result


def _write_csv(path: Path, rows: Sequence[Dict[str, Any]]) -> None:
    headers = _headers(rows)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _style_sheet(ws: Any) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape" if ws.max_column > 7 else "portrait"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.sheet_view.zoomScale = 90


def _write_table_sheet(wb: Any, title: str, rows: Sequence[Dict[str, Any]]) -> Any:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    headers = _headers(rows)
    width = max(1, len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
    cell.font = Font(name=FONT_NAME, size=15, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    thin = Side(style="thin", color="D9E1E8")
    for column, header in enumerate(headers, 1):
        cell = ws.cell(2, column, header)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[2].height = 26

    for row_index, row in enumerate(rows, 3):
        for column, header in enumerate(headers, 1):
            value = row.get(header)
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
            cell.alignment = Alignment(
                horizontal="right" if isinstance(value, (int, float)) and not isinstance(value, bool) else "left",
                vertical="center",
                wrap_text=False,
            )
            if isinstance(value, float):
                cell.number_format = "0.000000%" if header == "数值" and row.get("单位") == "%" else "0.00000000"
            if row_index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=LIGHT_FILL)
            cell.border = Border(bottom=thin)

    for column, header in enumerate(headers, 1):
        values = [str(header)] + [format_number(row.get(header)) for row in rows]
        longest = max((len(item) for item in values), default=8)
        is_numeric = any(isinstance(row.get(header), (int, float)) for row in rows)
        ws.column_dimensions[get_column_letter(column)].width = min(max(12 if is_numeric else 14, longest * 1.15 + 2), 30)
    _style_sheet(ws)
    return ws


def _write_summary_sheet(wb: Any, result: Dict[str, Any]) -> Any:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ws = wb.create_sheet("结果总览", 0)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "变焦镜头一阶设计计算结果"
    ws["A1"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    ws["A1"].font = Font(name=FONT_NAME, size=16, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    ws["A3"] = "方案"
    ws["B3"] = result["方案"]
    ws["C3"] = "结构模型"
    ws["D3"] = result["模型"]
    ws["A4"] = "综合状态"
    ws["B4"] = result["综合状态"]
    ws["C4"] = "生成时间"
    ws["D4"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    thin = Side(style="thin", color="D9E1E8")
    for row in ws.iter_rows(min_row=3, max_row=4, min_col=1, max_col=4):
        for cell in row:
            cell.font = Font(name=FONT_NAME, size=10, bold=cell.column in (1, 3), color=TEXT_COLOR)
            cell.fill = PatternFill("solid", fgColor=SECTION_FILL if cell.column in (1, 3) else ACCENT_FILL)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

    ws.merge_cells("A6:D6")
    ws["A6"] = "关键结果"
    ws["A6"].fill = PatternFill("solid", fgColor=SECTION_FILL)
    ws["A6"].font = Font(name=FONT_NAME, size=12, bold=True, color=TEXT_COLOR)

    headers = ["项目", "数值", "单位", "说明"]
    for column, header in enumerate(headers, 1):
        cell = ws.cell(7, column, header)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, item in enumerate(result["关键结果"], 8):
        for column, header in enumerate(headers, 1):
            value = item.get(header)
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
            cell.alignment = Alignment(horizontal="right" if column == 2 and isinstance(value, (int, float)) else "left", vertical="center")
            cell.border = Border(bottom=thin)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_FILL)
            if isinstance(value, float):
                cell.number_format = "0.000000%" if header == "数值" and item.get("单位") == "%" else "0.00000000"

    note_row = 9 + len(result["关键结果"])
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    ws.cell(note_row, 1, "适用边界")
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=WARNING_FILL)
    ws.cell(note_row, 1).font = Font(name=FONT_NAME, size=11, bold=True, color=TEXT_COLOR)
    for offset, note in enumerate(result.get("说明", []), 1):
        ws.merge_cells(start_row=note_row + offset, start_column=1, end_row=note_row + offset, end_column=4)
        cell = ws.cell(note_row + offset, 1, f"• {note}")
        cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row + offset].height = 24

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 46
    ws.freeze_panes = "A7"
    ws.sheet_view.zoomScale = 95
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    return ws


def _add_motion_chart(ws: Any, rows: Sequence[Dict[str, Any]]) -> None:
    from openpyxl.chart import LineChart, Reference

    headers = _headers(rows)
    if not rows or "位置" not in headers:
        return
    preferred = ["Delta2", "Delta3_推荐", "Delta3_连续"]
    series_headers = [header for header in preferred if header in headers]
    if not series_headers:
        return

    chart = LineChart()
    chart.title = "变倍组与补偿组位移"
    chart.style = 13
    chart.height = 10
    chart.width = 20
    chart.y_axis.title = "归一位移"
    chart.x_axis.title = "行程位置"
    category_column = headers.index("位置") + 1
    categories = Reference(ws, min_col=category_column, min_row=3, max_row=2 + len(rows))
    for header in series_headers:
        column = headers.index(header) + 1
        data = Reference(ws, min_col=column, min_row=2, max_row=2 + len(rows))
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend.position = "b"
    ws.add_chart(chart, f"{chr(65 + min(len(headers) + 1, 25))}2")


def _write_workbook(path: Path, result: Dict[str, Any]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _write_summary_sheet(wb, result)
    _write_table_sheet(wb, "输入参数", result["输入参数"])
    _write_table_sheet(wb, "组焦距", result["组焦距"])
    _write_table_sheet(wb, "长中短焦", result["长中短焦"])
    motion = _write_table_sheet(wb, "运动表", result["运动表"])
    _add_motion_chart(motion, result["运动表"])
    if result.get("物理结构"):
        _write_table_sheet(wb, "物理结构", result["物理结构"])
    if result.get("m3扫描比较"):
        _write_table_sheet(wb, "m3扫描比较", result["m3扫描比较"])
    wb.save(path)


def _pick(row: Dict[str, Any], names: Iterable[str]) -> Any:
    for name in names:
        if name in row:
            return row[name]
    return None


def _write_plot(path: Path, result: Dict[str, Any]) -> None:
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import pyplot as plt

    matplotlib.rcParams["font.sans-serif"] = ["Noto Sans CJK JP", "Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    matplotlib.rcParams["axes.unicode_minus"] = False
    rows = result["运动表"]
    positions = [100.0 * float(row["位置"]) for row in rows]
    delta2 = [float(row["Delta2"]) for row in rows]
    delta3 = [float(_pick(row, ("Delta3_推荐", "Delta3_连续", "Delta3"))) for row in rows]
    zoom = [float(_pick(row, ("倍率_推荐", "相对倍率"))) for row in rows]

    figure, axes = plt.subplots(2, 1, figsize=(10.5, 7.2), sharex=True, constrained_layout=True)
    figure.patch.set_facecolor("#F7F9FB")
    axes[0].plot(positions, delta2, color="#1F77B4", linewidth=2.2, marker="o", markersize=3.5, label="变倍组位移 Δ2")
    axes[0].plot(positions, delta3, color="#D97706", linewidth=2.2, marker="o", markersize=3.5, label="补偿组位移 Δ3")
    axes[0].axhline(0, color="#98A2B3", linewidth=0.8)
    axes[0].set_ylabel("归一位移")
    axes[0].set_title(f"{result['方案']}：组间运动")
    axes[0].legend(frameon=False, loc="best")

    axes[1].plot(positions, zoom, color="#2E7D32", linewidth=2.4, marker="o", markersize=3.5)
    axes[1].set_xlabel("变焦行程 (%)")
    axes[1].set_ylabel("相对倍率 (x)")
    axes[1].set_title("相对倍率随行程变化")
    for axis in axes:
        axis.set_facecolor("#FFFFFF")
        axis.grid(True, color="#D9E1E8", linewidth=0.7, alpha=0.8)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    figure.savefig(path, dpi=180, facecolor=figure.get_facecolor())
    plt.close(figure)


def _write_summary_text(path: Path, result: Dict[str, Any]) -> None:
    lines = [
        "变焦镜头一阶设计计算结果",
        "=" * 32,
        f"方案：{result['方案']}",
        f"模型：{result['模型']}",
        f"综合状态：{result['综合状态']}",
        "",
        "关键结果",
        "-" * 32,
    ]
    for item in result["关键结果"]:
        unit = f" {item.get('单位', '')}" if item.get("单位") else ""
        value = item.get("数值")
        shown = format_number(value * 100) if item.get("单位") == "%" and isinstance(value, (int, float)) else format_number(value)
        lines.append(f"{item['项目']}：{shown}{unit}")
    lines.extend(["", "组焦距", "-" * 32])
    for item in result["组焦距"]:
        normalized = format_number(item.get("归一焦距"))
        physical = item.get("物理焦距_mm")
        suffix = f"，{format_number(physical)} mm" if physical is not None else ""
        lines.append(f"{item['光学组']} ({item['符号']})：{normalized}{suffix}")
    if result.get("m3扫描比较"):
        lines.extend(["", "m3 扫描比较", "-" * 32])
        for row in result["m3扫描比较"]:
            lines.append(
                f"{row.get('m3M')} | {row.get('求解方式')} | {row.get('状态')}"
            )
    lines.extend(["", "说明", "-" * 32])
    lines.extend(f"- {note}" for note in result.get("说明", []))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def export_result(result: Dict[str, Any], base_dir: Path | str) -> Dict[str, Path]:
    output_dir = _make_output_dir(Path(base_dir), result["方案"])
    files = {
        "folder": output_dir,
        "xlsx": output_dir / "变焦一阶计算结果.xlsx",
        "csv": output_dir / "完整运动表.csv",
        "png": output_dir / "变焦运动曲线.png",
        "json": output_dir / "完整计算数据.json",
        "txt": output_dir / "结果摘要.txt",
    }
    _write_workbook(files["xlsx"], result)
    _write_csv(files["csv"], result["运动表"])
    if result.get("m3扫描比较"):
        files["m3_scan_csv"] = output_dir / "m3扫描比较.csv"
        _write_csv(files["m3_scan_csv"], result["m3扫描比较"])
    _write_plot(files["png"], result)
    files["json"].write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_summary_text(files["txt"], result)
    return files

